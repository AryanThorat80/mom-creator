from __future__ import annotations

import json
import logging
import re
import zipfile
from io import BytesIO
from typing import Any

import httpx

from app.core.config import settings

logger = logging.getLogger(__name__)


# ============================================================
# DEFAULT REPORT STRUCTURE
# ============================================================

DEFAULT_SECTIONS = [
    {
        "key": "meeting_overview",
        "title": "MEETING OVERVIEW",
        "source": "meeting_overview",
    },
    {
        "key": "participants",
        "title": "PARTICIPANTS",
        "source": "participants",
    },
    {
        "key": "agenda",
        "title": "MEETING AGENDA",
        "source": "agenda",
    },
    {
        "key": "summary",
        "title": "EXECUTIVE SUMMARY",
        "source": "summary",
    },
    {
        "key": "discussion",
        "title": "KEY DISCUSSION POINTS",
        "source": "discussion",
    },
    {
        "key": "decisions",
        "title": "DECISIONS AGREED UPON",
        "source": "decisions",
    },
    {
        "key": "action_items",
        "title": "ACTION ITEMS & OWNERSHIP MATRIX",
        "source": "action_items",
    },
    {
        "key": "risks",
        "title": "IDENTIFIED RISKS & DEPENDENCIES",
        "source": "risks",
    },
    {
        "key": "next_meeting",
        "title": "NEXT MEETING DETAILS",
        "source": "next_meeting",
    },
    {
        "key": "signoff",
        "title": "DOCUMENT SIGN-OFF & APPROVAL",
        "source": "signoff",
    },
]


# ============================================================
# REFERENCE TEXT EXTRACTION
# ============================================================

def _extract_pdf_text(file_bytes: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(file_bytes))

    pages: list[str] = []

    for page in reader.pages:
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""

        if text.strip():
            pages.append(text)

    return "\n\n".join(pages).strip()


def _extract_docx_text(file_bytes: bytes) -> str:
    try:
        from docx import Document
    except ImportError:
        return ""

    document = Document(BytesIO(file_bytes))

    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = []

            for cell in row.cells:
                cells.append(cell.text.strip())

            line = " | ".join(
                value for value in cells if value
            )

            if line:
                parts.append(line)

    return "\n".join(parts).strip()


def _extract_zip_xml_text(file_bytes: bytes) -> str:
    """
    Fallback DOCX extractor without requiring python-docx.
    """

    try:
        with zipfile.ZipFile(BytesIO(file_bytes)) as archive:
            xml = archive.read(
                "word/document.xml"
            ).decode(
                "utf-8",
                errors="ignore",
            )
    except Exception:
        return ""

    text = re.sub(
        r"<[^>]+>",
        " ",
        xml,
    )

    text = re.sub(
        r"\s+",
        " ",
        text,
    )

    return text.strip()


def extract_reference_text(file_bytes: bytes) -> str:
    """
    Extract text from PDF/DOCX reference files.

    Visual formatting is intentionally NOT copied.
    The reference is used to determine the report structure
    and required features.
    """

    if not file_bytes:
        return ""

    # PDF
    if file_bytes.startswith(b"%PDF"):
        return _extract_pdf_text(file_bytes)

    # DOCX
    if file_bytes.startswith(b"PK"):
        text = _extract_docx_text(file_bytes)

        if text:
            return text

        return _extract_zip_xml_text(file_bytes)

    # Plain text fallback
    try:
        return file_bytes.decode(
            "utf-8",
            errors="ignore",
        ).strip()
    except Exception:
        return ""


# ============================================================
# NORMALIZATION
# ============================================================

def _normalize_key(value: str) -> str:
    value = value.lower().strip()

    value = re.sub(
        r"[^a-z0-9]+",
        "_",
        value,
    )

    return value.strip("_")


def _fallback_reference_structure(
    reference_text: str,
) -> dict[str, Any]:

    text = reference_text.lower()

    sections: list[dict[str, str]] = []

    def add(
        key: str,
        title: str,
        source: str,
    ) -> None:

        sections.append(
            {
                "key": key,
                "title": title,
                "source": source,
            }
        )

    # Meeting overview is almost always present.
    add(
        "meeting_overview",
        "MEETING OVERVIEW",
        "meeting_overview",
    )

    if (
        "participant" in text
        or "attendee" in text
        or "attendance" in text
    ):
        add(
            "participants",
            "PARTICIPANTS",
            "participants",
        )

    if (
        "agenda" in text
        or "topic" in text
    ):
        add(
            "agenda",
            "MEETING AGENDA",
            "agenda",
        )

    if (
        "executive summary" in text
        or "summary" in text
    ):
        add(
            "summary",
            "EXECUTIVE SUMMARY",
            "summary",
        )

    if (
        "discussion" in text
        or "discussion points" in text
    ):
        add(
            "discussion",
            "KEY DISCUSSION POINTS",
            "discussion",
        )

    if (
        "decision" in text
        or "decisions agreed" in text
    ):
        add(
            "decisions",
            "DECISIONS AGREED UPON",
            "decisions",
        )

    if (
        "action item" in text
        or "ownership matrix" in text
        or "responsibilit" in text
    ):
        add(
            "action_items",
            "ACTION ITEMS & OWNERSHIP MATRIX",
            "action_items",
        )

    if (
        "risk" in text
        or "dependency" in text
    ):
        add(
            "risks",
            "IDENTIFIED RISKS & DEPENDENCIES",
            "risks",
        )

    if (
        "next meeting" in text
        or "proposed date" in text
        or "proposed time" in text
    ):
        add(
            "next_meeting",
            "NEXT MEETING DETAILS",
            "next_meeting",
        )

    if (
        "sign-off" in text
        or "approval" in text
        or "prepared by" in text
    ):
        add(
            "signoff",
            "DOCUMENT SIGN-OFF & APPROVAL",
            "signoff",
        )

    if not sections:
        sections = list(DEFAULT_SECTIONS)

    return {
        "sections": sections,
        "include_participants": any(
            section["key"] == "participants"
            for section in sections
        ),
    }


# ============================================================
# GEMINI REFERENCE ANALYSIS
# ============================================================

def _build_analysis_prompt(
    reference_text: str,
) -> str:

    return f"""
You are analyzing a Minutes of Meeting document template/example.

Your job is NOT to generate meeting content.

Your job is to identify the STRUCTURE and FEATURES that the final
MOM export must contain.

The uploaded document is authoritative for the report structure.

IMPORTANT:

1. Preserve every meaningful section/feature in the reference.
2. Do not remove a section simply because the current MOM does not
   contain information for it.
3. If the current MOM has no value for a section, the exporter will
   render that section as null/blank.
4. Preserve the logical order of the reference sections.
5. Detect participant/attendee sections explicitly.
6. Participants should ONLY be exported when the reference contains
   a participant/attendee/attendance section.
7. Do not invent sections that are not present.
8. Do not generate meeting content.
9. Do not copy example data from the reference.
10. Return valid JSON only.

Return exactly:

{{
  "include_participants": true,
  "sections": [
    {{
      "key": "meeting_overview",
      "title": "MEETING OVERVIEW",
      "source": "meeting_overview"
    }}
  ]
}}

Allowed source values:

meeting_overview
participants
agenda
summary
discussion
decisions
action_items
risks
next_meeting
signoff
abbreviations
transcript
custom

For custom sections use:

"source": "custom"

and preserve the actual section title.

REFERENCE DOCUMENT:

-------------------------
{reference_text[:50000]}
-------------------------
"""


async def _call_gemini(
    prompt: str,
) -> dict[str, Any]:

    api_key = getattr(
        settings,
        "ai_api_key",
        None,
    )

    base_url = getattr(
        settings,
        "ai_base_url",
        None,
    )

    model = getattr(
        settings,
        "ai_model",
        None,
    )

    if not api_key:
        raise RuntimeError(
            "AI API key is not configured"
        )

    if not base_url:
        base_url = (
            "https://generativelanguage.googleapis.com"
        )

    if not model:
        model = "gemini-2.5-flash"

    base_url = base_url.rstrip("/")

    url = (
        f"{base_url}/v1beta/models/"
        f"{model}:generateContent"
    )

    payload = {
        "contents": [
            {
                "parts": [
                    {
                        "text": prompt,
                    }
                ]
            }
        ],
        "generationConfig": {
            "temperature": 0,
            "responseMimeType": "application/json",
        },
    }

    headers = {
        "x-goog-api-key": api_key,
        "Content-Type": "application/json",
    }

    async with httpx.AsyncClient(
        timeout=90,
    ) as client:

        response = await client.post(
            url,
            headers=headers,
            json=payload,
        )

    if response.status_code >= 400:
        raise RuntimeError(
            "Gemini reference analysis failed: "
            f"HTTP {response.status_code}: "
            f"{response.text[:1000]}"
        )

    data = response.json()

    candidates = data.get(
        "candidates",
        [],
    )

    if not candidates:
        raise RuntimeError(
            "Gemini returned no reference analysis"
        )

    parts = (
        candidates[0]
        .get("content", {})
        .get("parts", [])
    )

    generated = ""

    for part in parts:
        value = part.get("text")

        if value:
            generated += value

    generated = generated.strip()

    if not generated:
        raise RuntimeError(
            "Gemini returned empty reference analysis"
        )

    return json.loads(generated)


# ============================================================
# PUBLIC REFERENCE ANALYZER
# ============================================================

async def analyze_reference(
    file_bytes: bytes,
) -> dict[str, Any]:

    reference_text = extract_reference_text(
        file_bytes
    )

    if not reference_text:
        return _fallback_reference_structure(
            ""
        )

    try:

        result = await _call_gemini(
            _build_analysis_prompt(
                reference_text
            )
        )

        sections = result.get(
            "sections",
            [],
        )

        if not isinstance(
            sections,
            list,
        ) or not sections:

            raise RuntimeError(
                "Invalid section structure"
            )

        cleaned_sections = []

        for section in sections:

            if not isinstance(
                section,
                dict,
            ):
                continue

            title = str(
                section.get(
                    "title",
                    "",
                )
            ).strip()

            if not title:
                continue

            key = _normalize_key(
                str(
                    section.get(
                        "key",
                        title,
                    )
                )
            )

            source = str(
                section.get(
                    "source",
                    "custom",
                )
            ).strip()

            cleaned_sections.append(
                {
                    "key": key,
                    "title": title,
                    "source": source,
                }
            )

        if not cleaned_sections:
            raise RuntimeError(
                "Reference contained no usable sections"
            )

        include_participants = bool(
            result.get(
                "include_participants",
                any(
                    section["source"]
                    == "participants"
                    for section in cleaned_sections
                ),
            )
        )

        return {
            "sections": cleaned_sections,
            "include_participants": (
                include_participants
            ),
        }

    except Exception as exc:

        logger.warning(
            "Reference AI analysis failed. "
            "Using structural fallback: %s",
            exc,
        )

        return _fallback_reference_structure(
            reference_text
        )