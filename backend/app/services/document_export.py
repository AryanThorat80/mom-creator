import io
import re
import zipfile
from datetime import datetime, timezone, timedelta
from html import escape
from typing import Any
from xml.sax.saxutils import escape as xml_escape

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


IST = timezone(timedelta(hours=5, minutes=30))
XL_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def safe_filename(value: str | None, fallback: str = "mom") -> str:
    value = (value or "").strip()
    value = re.sub(r"[^\w\s-]", "", value, flags=re.UNICODE)
    value = re.sub(r"[\s_-]+", "-", value).strip("-")
    return value[:80] or fallback


def format_datetime(value: Any) -> str:
    if not value:
        return "Not specified"
    if isinstance(value, datetime):
        dt = value
    else:
        try:
            dt = datetime.fromisoformat(str(value).strip().replace("Z", "+00:00"))
        except ValueError:
            return str(value)
    if dt.tzinfo is not None:
        dt = dt.astimezone(IST)
    return dt.strftime("%d %b %Y, %I:%M %p").lstrip("0")


def format_date_only(value: Any) -> str:
    if not value:
        return "No due date"
    if isinstance(value, datetime):
        dt = value
    else:
        try:
            dt = datetime.fromisoformat(str(value).strip().replace("Z", "+00:00"))
        except ValueError:
            return str(value)
    if dt.tzinfo is not None:
        dt = dt.astimezone(IST)
    return dt.strftime("%d %b %Y")


def normalize_list(value: Any) -> list[str]:
    return [str(x).strip() for x in value if str(x).strip()] if isinstance(value, list) else []


def build_export_payload(meeting, mom, action_items):
    return {"meeting": meeting, "mom": mom, "action_items": action_items}


# ============================================================
# DOCX
# ============================================================


def _shade_cell(cell, fill: str = "E2E8F0") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_run_font(run, name: str = "Aptos", size: float = 10.5, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)


def _setup_docx(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Aptos")
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Aptos")

    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)


def _add_bullets(doc: Document, items: list[str]) -> None:
    if not items:
        doc.add_paragraph("None recorded.")
        return
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        _set_run_font(p.add_run(item))


def _add_action_table(doc: Document, action_items: list[dict[str, Any]]) -> None:
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["Task", "Assignee", "Due Date", "Status", "Priority"]
    for index, label in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = label
        _shade_cell(cell)
        for run in cell.paragraphs[0].runs:
            _set_run_font(run, size=9, bold=True)

    if not action_items:
        row = table.add_row().cells
        row[0].merge(row[-1])
        row[0].text = "No action items recorded."
        return

    for item in action_items:
        row = table.add_row().cells
        values = [
            item.get("task") or "",
            item.get("assigned_name") or "Unassigned",
            format_date_only(item.get("due_date")),
            item.get("status") or "pending",
            item.get("priority") or "medium",
        ]
        for i, value in enumerate(values):
            row[i].text = str(value)
            row[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for paragraph in row[i].paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    _set_run_font(run, size=8.5)


def generate_docx(meeting: dict[str, Any], mom: dict[str, Any], action_items: list[dict[str, Any]]) -> bytes:
    doc = Document()
    _setup_docx(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(
        title.add_run(mom.get("title") or meeting.get("title") or "Minutes of Meeting"),
        name="Aptos Display",
        size=22,
        bold=True,
    )

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(subtitle.add_run("MOM Creator"), size=10, bold=True)

    metadata = doc.add_table(rows=3, cols=2)
    metadata.style = "Table Grid"
    meta = [
        ("Meeting Date", format_datetime(meeting.get("meeting_date"))),
        ("Meeting Mode", meeting.get("mode") or "Not specified"),
        ("MOM Status", mom.get("status") or "draft"),
    ]
    for r, (label, value) in enumerate(meta):
        metadata.cell(r, 0).text = label
        metadata.cell(r, 1).text = str(value)
        _shade_cell(metadata.cell(r, 0), "F1F5F9")
        for run in metadata.cell(r, 0).paragraphs[0].runs:
            _set_run_font(run, bold=True)

    doc.add_paragraph()

    doc.add_heading("Summary", level=1)
    summary_paragraph = doc.add_paragraph()
    _set_run_font(summary_paragraph.add_run(mom.get("summary") or "No summary recorded."))

    doc.add_heading("Key Discussion Points", level=1)
    _add_bullets(doc, normalize_list(mom.get("key_discussion_points")))

    doc.add_heading("Decisions", level=1)
    _add_bullets(doc, normalize_list(mom.get("decisions")))

    doc.add_heading("Next Steps", level=1)
    _add_bullets(doc, normalize_list(mom.get("next_steps")))

    doc.add_heading("Action Items", level=1)
    _add_action_table(doc, action_items)

    abbreviations = mom.get("abbreviations_used") or {}
    if abbreviations:
        doc.add_heading("Abbreviations", level=1)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        for i, label in enumerate(["Abbreviation", "Meaning"]):
            table.rows[0].cells[i].text = label
            _shade_cell(table.rows[0].cells[i])
            for run in table.rows[0].cells[i].paragraphs[0].runs:
                _set_run_font(run, size=9, bold=True)
        for key, meaning in abbreviations.items():
            row = table.add_row().cells
            row[0].text = str(key)
            row[1].text = str(meaning)

    transcript = (mom.get("transcript") or "").strip()
    if transcript:
        doc.add_page_break()
        doc.add_heading("Transcript", level=1)
        p = doc.add_paragraph()
        _set_run_font(p.add_run(transcript), size=9)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(footer.add_run("Generated by MOM Creator"), size=8)

    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


# ============================================================
# XLSX, written as OOXML directly so the backend has no
# dependency on openpyxl/xlsxwriter/pandas.
# ============================================================


def _column_letter(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _xlsx_cell(row: int, col: int, value: Any, style: int = 0) -> str:
    ref = f"{_column_letter(col)}{row}"
    if value is None:
        return f'<c r="{ref}" s="{style}"/>'
    return (
        f'<c r="{ref}" s="{style}" t="inlineStr">'
        f'<is><t xml:space="preserve">{xml_escape(str(value))}</t></is></c>'
    )


def _xlsx_sheet_xml(rows, widths):
    cols = "".join(
        f'<col min="{i}" max="{i}" width="{width}" customWidth="1"/>'
        for i, width in enumerate(widths, start=1)
    )
    row_xml = ""
    for row_num, row in enumerate(rows, start=1):
        cells = "".join(
            _xlsx_cell(row_num, col_num, value, 1 if row_num == 1 else 0)
            for col_num, value in enumerate(row, start=1)
        )
        row_xml += f'<row r="{row_num}">{cells}</row>'
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<worksheet xmlns="{XL_NS}"><cols>{cols}</cols><sheetData>{row_xml}</sheetData></worksheet>'
    )


def _xlsx_workbook_xml(names):
    sheets = "".join(
        f'<sheet name="{xml_escape(name)}" sheetId="{i}" r:id="rId{i}"/>'
        for i, name in enumerate(names, start=1)
    )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<workbook xmlns="{XL_NS}" xmlns:r="{REL_NS}"><sheets>{sheets}</sheets></workbook>'
    )


def _xlsx_rels(count):
    rels = "".join(
        f'<Relationship Id="rId{i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{i}.xml"/>'
        for i in range(1, count + 1)
    )
    rels += (
        '<Relationship Id="rIdStyles" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
    )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f"{rels}</Relationships>"
    )


def _xlsx_styles():
    return '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <fonts count="2">
    <font><sz val="10"/><name val="Aptos"/></font>
    <font><b/><sz val="10"/><name val="Aptos"/></font>
  </fonts>
  <fills count="2">
    <fill><patternFill patternType="none"/></fill>
    <fill><patternFill patternType="solid"><fgColor rgb="E2E8F0"/><bgColor indexed="64"/></patternFill></fill>
  </fills>
  <borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>
  <cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>
  <cellXfs count="2">
    <xf numFmtId="0" fontId="0" fillId="0" borderId="0" applyAlignment="1"><alignment vertical="top" wrapText="1"/></xf>
    <xf numFmtId="0" fontId="1" fillId="1" borderId="0" applyAlignment="1"><alignment vertical="center" wrapText="1"/></xf>
  </cellXfs>
  <cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>
</styleSheet>'''


def _xlsx_content_types(count):
    overrides = "".join(
        f'<Override PartName="/xl/worksheets/sheet{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        for i in range(1, count + 1)
    )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
        f"{overrides}</Types>"
    )


def generate_xlsx(meeting: dict[str, Any], mom: dict[str, Any], action_items: list[dict[str, Any]]) -> bytes:
    sheets = [
        (
            "Summary",
            [
                ["Field", "Value"],
                ["Meeting Title", mom.get("title") or meeting.get("title") or ""],
                ["Meeting Date", format_datetime(meeting.get("meeting_date"))],
                ["Meeting Mode", meeting.get("mode") or ""],
                ["MOM Status", mom.get("status") or ""],
                ["Summary", mom.get("summary") or ""],
            ],
            [24, 100],
        ),
        (
            "Discussion Points",
            [["#", "Discussion Point"]]
            + [[i, value] for i, value in enumerate(normalize_list(mom.get("key_discussion_points")), 1)],
            [8, 110],
        ),
        (
            "Decisions",
            [["#", "Decision"]]
            + [[i, value] for i, value in enumerate(normalize_list(mom.get("decisions")), 1)],
            [8, 110],
        ),
        (
            "Action Items",
            [["ID", "Task", "Assignee", "Due Date", "Status", "Priority", "Google Event ID"]]
            + [
                [
                    item.get("id") or "",
                    item.get("task") or "",
                    item.get("assigned_name") or "Unassigned",
                    format_date_only(item.get("due_date")),
                    item.get("status") or "pending",
                    item.get("priority") or "medium",
                    item.get("google_event_id") or "",
                ]
                for item in action_items
            ],
            [38, 75, 22, 18, 16, 12, 30],
        ),
        (
            "Abbreviations",
            [["Abbreviation", "Meaning"]]
            + [[key, value] for key, value in (mom.get("abbreviations_used") or {}).items()],
            [20, 80],
        ),
        (
            "Transcript",
            [["Transcript"], [mom.get("transcript") or "No transcript recorded."]],
            [120],
        ),
    ]

    names = [name for name, _, _ in sheets]
    package_rels = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rIdWorkbook" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>'''

    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", _xlsx_content_types(len(sheets)))
        archive.writestr("_rels/.rels", package_rels)
        archive.writestr("xl/workbook.xml", _xlsx_workbook_xml(names))
        archive.writestr("xl/_rels/workbook.xml.rels", _xlsx_rels(len(sheets)))
        archive.writestr("xl/styles.xml", _xlsx_styles())
        for i, (_, rows, widths) in enumerate(sheets, start=1):
            archive.writestr(f"xl/worksheets/sheet{i}.xml", _xlsx_sheet_xml(rows, widths))
    return stream.getvalue()


# ============================================================
# PRINTABLE HTML
# ============================================================

PRINT_CSS = """
@page { size: A4; margin: 16mm; }
* { box-sizing: border-box; }
body { font-family: Arial, Helvetica, sans-serif; color: #111827; margin: 0; line-height: 1.5; font-size: 13px; }
.page { max-width: 850px; margin: 0 auto; }
.header { text-align: center; margin-bottom: 24px; }
.header h1 { margin: 0; font-size: 26px; }
.header p { margin: 6px 0 0; color: #64748b; }
.meta { width: 100%; border-collapse: collapse; margin-bottom: 24px; }
.meta th, .meta td { border: 1px solid #cbd5e1; padding: 8px 10px; text-align: left; }
.meta th { width: 180px; background: #f1f5f9; }
section { margin-top: 20px; }
h2 { font-size: 17px; border-bottom: 1px solid #cbd5e1; padding-bottom: 5px; }
.action { width: 100%; border-collapse: collapse; }
.action th, .action td { border: 1px solid #cbd5e1; padding: 7px; vertical-align: top; text-align: left; }
.action th { background: #e2e8f0; }
.transcript { white-space: pre-wrap; background: #f8fafc; border: 1px solid #e2e8f0; padding: 12px; }
@media print { .page { max-width: none; } .avoid-break { break-inside: avoid; } }
"""


def generate_printable_html(meeting, mom, action_items) -> str:
    title = escape(mom.get("title") or meeting.get("title") or "Minutes of Meeting")
    summary = escape(mom.get("summary") or "No summary recorded.")

    def bullets(items):
        if not items:
            return "<p>None recorded.</p>"
        return "<ul>" + "".join(f"<li>{escape(x)}</li>" for x in items) + "</ul>"

    rows = "".join(
        "<tr>"
        f"<td>{escape(str(item.get('task') or ''))}</td>"
        f"<td>{escape(str(item.get('assigned_name') or 'Unassigned'))}</td>"
        f"<td>{escape(format_date_only(item.get('due_date')))}</td>"
        f"<td>{escape(str(item.get('status') or 'pending'))}</td>"
        f"<td>{escape(str(item.get('priority') or 'medium'))}</td>"
        "</tr>"
        for item in action_items
    )
    if not rows:
        rows = '<tr><td colspan="5">No action items recorded.</td></tr>'

    transcript = escape((mom.get("transcript") or "").strip() or "No transcript recorded.")

    return f'''<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>{title}</title>
<style>{PRINT_CSS}</style>
</head>
<body>
<main class="page">
<header class="header"><h1>{title}</h1><p>MOM Creator</p></header>
<table class="meta">
<tr><th>Meeting Date</th><td>{escape(format_datetime(meeting.get("meeting_date")))}</td></tr>
<tr><th>Meeting Mode</th><td>{escape(str(meeting.get("mode") or "Not specified"))}</td></tr>
<tr><th>MOM Status</th><td>{escape(str(mom.get("status") or "draft"))}</td></tr>
</table>
<section class="avoid-break"><h2>Summary</h2><p>{summary}</p></section>
<section class="avoid-break"><h2>Key Discussion Points</h2>{bullets(normalize_list(mom.get("key_discussion_points")))}</section>
<section class="avoid-break"><h2>Decisions</h2>{bullets(normalize_list(mom.get("decisions")))}</section>
<section class="avoid-break"><h2>Next Steps</h2>{bullets(normalize_list(mom.get("next_steps")))}</section>
<section><h2>Action Items</h2>
<table class="action"><thead><tr><th>Task</th><th>Assignee</th><th>Due Date</th><th>Status</th><th>Priority</th></tr></thead>
<tbody>{rows}</tbody></table></section>
<section><h2>Transcript</h2><div class="transcript">{transcript}</div></section>
</main>
</body>
</html>'''
